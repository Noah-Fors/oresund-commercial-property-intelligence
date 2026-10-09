from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from src.phase3_market_report.data import (
    LISTINGS_SHEET,
    METRICS_SHEET,
    TEMPLATE_VERSION,
    create_template,
    load_listings,
    load_metrics,
    upgrade_workbook,
)
from src.phase3_market_report.models import MarketArea, MarketMetric, Metric, OfficeListing, Segment


def _fill(path, sheet, rows):
    wb = load_workbook(path)
    ws = wb[sheet]
    header = [cell.value for cell in ws[1]]
    for r, row in enumerate(rows, start=2):
        for key, value in row.items():
            ws.cell(row=r, column=header.index(key) + 1, value=value)
    wb.save(path)


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "market.xlsx"
    create_template(path)
    return path


def test_template_is_never_overwritten(workbook):
    with pytest.raises(FileExistsError):
        create_template(workbook)


def test_empty_template_loads_without_rows_or_errors(workbook):
    assert load_metrics(workbook).rows == []
    assert load_listings(workbook).errors == []


def test_single_value_metric_gets_high_equal_to_low():
    metric = MarketMetric(city="malmo", submarket="CBD", segment="office", metric="prime_yield", low=4.85,
                          source="Exempel", published=date(2026, 3, 31))
    assert metric.high == 4.85
    assert metric.unit == "%"


def test_percent_written_as_decimal_is_rejected():
    with pytest.raises(ValueError, match="4.85"):
        MarketMetric(city="malmo", submarket="CBD", segment="office", metric="vacancy_rate", low=0.10,
                     source="Exempel", published=date(2026, 3, 31))


def test_rent_written_as_yield_is_rejected():
    # Prime yield 4,52 % inmatat som decimaltal, med prime_rent valt i rullistan.
    with pytest.raises(ValueError, match="prime_rent"):
        MarketMetric(city="malmo", submarket="CBD", segment="office", metric="prime_rent", low=0.0452,
                     source="Exempel", published=date(2026, 6, 30))


def test_implausible_low_end_of_range_is_rejected():
    with pytest.raises(ValueError, match="rent_range"):
        MarketMetric(city="lund", submarket="Lägesklass A", segment="office", metric="rent_range", low=170, high=2700,
                     source="Exempel", published=date(2026, 6, 30))


def test_monthly_rent_is_caught():
    with pytest.raises(ValueError, match="per månad"):
        OfficeListing(address="Stortorget 1", city="malmo", district="Centrum", area_sqm=200,
                      asking_rent_per_sqm=35_000, source="Exempel", collected=date(2026, 10, 6))


def test_metrics_sheet_reads_excel_dates_and_ranges(workbook):
    _fill(workbook, METRICS_SHEET, [
        {"city": "lund", "submarket": "Lägesklass A", "segment": "industrial", "metric": "rent_range", "low": 1700, "high": 2700,
         "source": "Exempel", "published": datetime(2026, 6, 30)},
    ])
    result = load_metrics(workbook)
    assert result.errors == []
    [metric] = result.rows
    assert metric.metric is Metric.RENT_RANGE
    assert metric.segment is Segment.INDUSTRIAL
    assert metric.published == date(2026, 6, 30)
    assert metric.midpoint == 2200


def test_listing_errors_point_to_the_excel_row(workbook):
    valid = {"address": "Stortorget 1", "city": "malmo", "district": "Centrum", "area_sqm": 250,
             "asking_rent_per_sqm": 2100, "source": "Exempel", "collected": "2026-10-06",
             "latitude": 55.6066, "longitude": 13.0010}
    _fill(workbook, LISTINGS_SHEET, [
        valid,
        {**valid, "city": "Malmö"},
        {**valid, "longitude": None},
        {**valid, "asking_rent_per_sqm": None},
    ])
    result = load_listings(workbook)
    assert len(result.rows) == 2
    assert result.rows[0].coordinates.latitude == 55.6066
    assert result.rows[1].asking_rent_per_sqm is None
    assert result.errors[0].startswith("rad 3: city")
    assert result.errors[1].startswith("rad 4: coordinates")


def test_segment_is_required(workbook):
    _fill(workbook, METRICS_SHEET, [
        {"city": "malmo", "submarket": "Lägesklass A", "metric": "rent_range", "low": 1600, "high": 3000,
         "source": "Exempel", "published": "2026-10-09"},
    ])
    assert load_metrics(workbook).errors == ["rad 2: segment: Input should be 'office', 'retail', 'industrial' or 'residential'"]


def test_upgrade_moves_rows_from_a_workbook_without_segment_and_marks_them_office(workbook):
    wb = load_workbook(workbook)
    wb[METRICS_SHEET].delete_cols(3)  # så såg mallen ut innan segment fanns
    info = wb["Instruktioner"]
    for row in info.iter_rows():  # version 1 hade inget versionsnummer
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("Mallversion"):
                cell.value = None
    wb.save(workbook)
    old_row = {"city": "malmo", "submarket": "Lägesklass AA", "metric": "rent_range", "low": 1800,
               "high": 3600, "source": "Newsec via Objektvision", "published": datetime(2026, 10, 9)}
    _fill(workbook, METRICS_SHEET, [old_row])
    _fill(workbook, LISTINGS_SHEET, [{"address": "Stortorget 1", "city": "malmo"}])

    backup = upgrade_workbook(workbook)

    assert backup.exists()
    [metric] = load_metrics(workbook).rows
    assert metric.segment is Segment.OFFICE
    assert (metric.low, metric.high, metric.published) == (1800, 3600, date(2026, 10, 9))
    assert load_workbook(workbook)[LISTINGS_SHEET]["A2"].value == "Stortorget 1"
    assert upgrade_workbook(workbook) is None  # andra gången finns inget att uppgradera


def test_regional_figures_are_allowed_for_market_metrics_but_not_listings():
    metric = MarketMetric(city="oresund", submarket="Öresund", segment="industrial", metric="prime_yield",
                          low=5.2, source="Exempel", published=date(2026, 6, 30))
    assert metric.city is MarketArea.ORESUND
    with pytest.raises(ValueError):
        OfficeListing(address="Stortorget 1", city="oresund", district="Centrum", area_sqm=200,
                      source="Exempel", collected=date(2026, 10, 9))


def test_residential_accepts_yield_but_not_rent():
    MarketMetric(city="malmo", submarket="Hela staden", segment="residential", metric="prime_yield",
                 low=4.0, source="Exempel", published=date(2026, 6, 30))
    with pytest.raises(ValueError, match="reglerade"):
        MarketMetric(city="malmo", submarket="Hela staden", segment="residential", metric="rent_range",
                     low=1200, high=1800, source="Exempel", published=date(2026, 6, 30))


def test_new_template_is_current_and_not_upgraded(workbook):
    from src.phase3_market_report.data import template_version
    assert template_version(load_workbook(workbook)) == TEMPLATE_VERSION
    assert upgrade_workbook(workbook) is None


def test_upgrade_from_version_2_gets_new_dropdowns_and_keeps_rows(workbook):
    wb = load_workbook(workbook)
    for row in wb["Instruktioner"].iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("Mallversion"):
                cell.value = "Mallversion: 2"
    wb.save(workbook)
    _fill(workbook, METRICS_SHEET, [
        {"city": "malmo", "submarket": "CBD", "segment": "office", "metric": "prime_yield", "low": 4.8,
         "source": "Exempel", "published": datetime(2026, 6, 30)},
    ])

    assert upgrade_workbook(workbook) is not None

    [metric] = load_metrics(workbook).rows
    assert metric.low == 4.8
    validations = load_workbook(workbook)[METRICS_SHEET].data_validations.dataValidation
    assert any("oresund" in dv.formula1 for dv in validations)
    assert any("residential" in dv.formula1 for dv in validations)
